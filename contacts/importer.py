import csv
import io
import openpyxl
from .utils import normalize_phone_number
from .models import Contact, ContactGroup, ContactGroupMembership


class ContactImporter:
    """
    Parses and imports contacts from CSV or XLSX files.
    Supports previewing, field mapping, validation, duplicate handling, and reporting.
    """

    ALLOWED_DUPLICATE_POLICIES = ['SKIP', 'UPDATE']

    def __init__(self, business, group_id=None, duplicate_policy='SKIP'):
        self.business = business
        self.group = None
        if group_id:
            self.group = ContactGroup.objects.filter(business=business, id=group_id).first()
        self.duplicate_policy = duplicate_policy if duplicate_policy in self.ALLOWED_DUPLICATE_POLICIES else 'SKIP'

    def parse_preview(self, uploaded_file, max_rows=5):
        """
        Reads first few rows to preview headers and sample data.
        Returns:
            dict: {
                'filename': str,
                'headers': list of str,
                'sample_rows': list of lists,
                'total_estimated_rows': int,
                'file_format': 'csv' | 'xlsx'
            }
        """
        filename = uploaded_file.name.lower()
        if filename.endswith('.csv'):
            return self._preview_csv(uploaded_file, max_rows)
        elif filename.endswith('.xlsx'):
            return self._preview_xlsx(uploaded_file, max_rows)
        else:
            raise ValueError("Unsupported file format. Please upload a CSV (.csv) or Excel (.xlsx) file.")

    def _preview_csv(self, file_obj, max_rows):
        content = file_obj.read()
        # Handle UTF-8 with BOM or plain UTF-8, Latin-1 fallback
        text = None
        for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            raise ValueError("Could not decode CSV file. Please ensure it is saved with UTF-8 encoding.")

        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            raise ValueError("The uploaded CSV file is empty.")

        headers = [h.strip() for h in rows[0]]
        sample_rows = rows[1:max_rows + 1]
        return {
            'filename': file_obj.name,
            'headers': headers,
            'sample_rows': sample_rows,
            'total_estimated_rows': max(0, len(rows) - 1),
            'file_format': 'csv'
        }

    def _preview_xlsx(self, file_obj, max_rows):
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        sheet = wb.active
        all_rows = list(sheet.iter_rows(values_only=True))
        if not all_rows:
            raise ValueError("The uploaded Excel sheet is empty.")

        headers = [str(h).strip() if h is not None else f"Column_{i+1}" for i, h in enumerate(all_rows[0])]
        sample_rows = [
            [str(cell) if cell is not None else '' for cell in row]
            for row in all_rows[1:max_rows + 1]
        ]
        return {
            'filename': file_obj.name,
            'headers': headers,
            'sample_rows': sample_rows,
            'total_estimated_rows': max(0, len(all_rows) - 1),
            'file_format': 'xlsx'
        }

    def process_import(self, uploaded_file, name_col_idx=0, phone_col_idx=1, email_col_idx=None):
        """
        Executes full validation and import into the database.
        Returns:
            dict: {
                'total_rows': int,
                'created_count': int,
                'updated_count': int,
                'skipped_count': int,
                'invalid_count': int,
                'errors': list of dicts [{'row': i, 'name': n, 'phone': p, 'error': msg}],
            }
        """
        filename = uploaded_file.name.lower()
        if filename.endswith('.csv'):
            raw_rows = self._read_csv_rows(uploaded_file)
        elif filename.endswith('.xlsx'):
            raw_rows = self._read_xlsx_rows(uploaded_file)
        else:
            raise ValueError("Unsupported file format.")

        if not raw_rows:
            return {
                'total_rows': 0, 'created_count': 0, 'updated_count': 0,
                'skipped_count': 0, 'invalid_count': 0, 'errors': []
            }

        headers = raw_rows[0]
        data_rows = raw_rows[1:]

        created_count = 0
        updated_count = 0
        skipped_count = 0
        invalid_count = 0
        errors = []

        seen_phones_in_batch = set()

        for idx, row in enumerate(data_rows, start=2):
            if not row or all(str(c).strip() == '' for c in row):
                continue  # skip completely empty lines

            # Extract Name
            name = ''
            if name_col_idx is not None and 0 <= name_col_idx < len(row):
                name = str(row[name_col_idx]).strip() if row[name_col_idx] is not None else ''

            # Extract Phone
            raw_phone = ''
            if phone_col_idx is not None and 0 <= phone_col_idx < len(row):
                raw_phone = str(row[phone_col_idx]).strip() if row[phone_col_idx] is not None else ''

            # Extract Email
            email = ''
            if email_col_idx is not None and 0 <= email_col_idx < len(row):
                email = str(row[email_col_idx]).strip() if row[email_col_idx] is not None else ''

            if not name:
                name = f"Contact {raw_phone}"

            # Validate Phone
            is_valid, normalized_phone, error_msg = normalize_phone_number(raw_phone)
            if not is_valid:
                invalid_count += 1
                errors.append({
                    'row': idx,
                    'name': name,
                    'phone': raw_phone,
                    'error': error_msg
                })
                continue

            # Check duplicate within the same batch file
            if normalized_phone in seen_phones_in_batch:
                skipped_count += 1
                errors.append({
                    'row': idx,
                    'name': name,
                    'phone': raw_phone,
                    'error': f"Duplicate phone number in uploaded file ({normalized_phone})"
                })
                continue

            seen_phones_in_batch.add(normalized_phone)

            # Check existing contact in database for this business
            existing_contact = Contact.objects.filter(business=self.business, phone_number=normalized_phone).first()

            if existing_contact:
                if self.duplicate_policy == 'UPDATE':
                    existing_contact.name = name or existing_contact.name
                    if email:
                        existing_contact.email = email
                    existing_contact.save()
                    updated_count += 1
                    target_contact = existing_contact
                else:  # 'SKIP'
                    skipped_count += 1
                    target_contact = existing_contact
            else:
                target_contact = Contact.objects.create(
                    business=self.business,
                    name=name,
                    phone_number=normalized_phone,
                    email=email or None,
                )
                created_count += 1

            # Attach to group if specified
            if self.group and target_contact:
                ContactGroupMembership.objects.get_or_create(
                    contact=target_contact,
                    group=self.group
                )

        return {
            'total_rows': len(data_rows),
            'created_count': created_count,
            'updated_count': updated_count,
            'skipped_count': skipped_count,
            'invalid_count': invalid_count,
            'errors': errors
        }

    def _read_csv_rows(self, file_obj):
        content = file_obj.read()
        for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        reader = csv.reader(io.StringIO(text))
        return list(reader)

    def _read_xlsx_rows(self, file_obj):
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        sheet = wb.active
        return list(sheet.iter_rows(values_only=True))
